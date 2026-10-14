import sqlite3
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

import app as backend


@pytest.fixture
def api(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    schema_path = Path(__file__).resolve().parents[1] / "schema.sql"
    conn = sqlite3.connect(db_path)
    conn.executescript(schema_path.read_text(encoding="utf-8"))
    conn.execute(
        "INSERT INTO users (name, email, password_hash, role) VALUES (?, ?, ?, ?)",
        ("Owner", "owner@test.local", generate_password_hash("owner-password"), "owner"),
    )
    conn.execute(
        "INSERT INTO users (name, email, password_hash, role) VALUES (?, ?, ?, ?)",
        ("Staff", "staff@test.local", generate_password_hash("staff-password"), "staff"),
    )
    conn.executemany(
        """INSERT INTO products
           (name, brand, category, season, min_stock_level, is_restricted)
           VALUES (?, ?, ?, ?, ?, ?)""",
        [
            ("General Product", "Test", "Fungicide", "Year-round", 10, 0),
            ("Monocrotophos", "Test", "Insecticide", "Kharif", 2, 1),
        ],
    )
    today = date.today()
    conn.executemany(
        """INSERT INTO batches
           (product_id, batch_no, quantity_received, quantity_remaining,
            price_per_unit, purchase_date, expiry_date)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        [
            (1, "OLD-EXPIRY", 3, 3, 10, today.isoformat(),
             (today + timedelta(days=10)).isoformat()),
            (1, "LATER-EXPIRY", 4, 4, 12, today.isoformat(),
             (today + timedelta(days=20)).isoformat()),
            (2, "RESTRICTED", 5, 5, 8, today.isoformat(),
             (today + timedelta(days=30)).isoformat()),
        ],
    )
    conn.commit()
    conn.close()

    monkeypatch.setattr(backend, "DB_NAME", str(db_path))
    backend.app.config.update(TESTING=True)
    client = backend.app.test_client()
    owner_login = client.post(
        "/auth/login",
        json={"email": "owner@test.local", "password": "owner-password"},
    )
    staff_login = client.post(
        "/auth/login",
        json={"email": "staff@test.local", "password": "staff-password"},
    )
    return {
        "client": client,
        "db_path": db_path,
        "owner": {"Authorization": f"Bearer {owner_login.json['token']}"},
        "staff": {"Authorization": f"Bearer {staff_login.json['token']}"},
    }


def test_fifo_sale_uses_earliest_expiry_first(api):
    response = api["client"].post(
        "/sales",
        json={"product_id": 1, "quantity_sold": 5},
        headers=api["staff"],
    )

    assert response.status_code == 201
    assert [batch["quantity_taken"] for batch in response.json["batches_used"]] == [3, 2]
    conn = sqlite3.connect(api["db_path"])
    remaining = conn.execute(
        "SELECT quantity_remaining FROM batches WHERE product_id = 1 ORDER BY expiry_date"
    ).fetchall()
    conn.close()
    assert remaining == [(0,), (2,)]


def test_restricted_sale_requires_buyer_id_and_appears_in_compliance_log(api):
    client = api["client"]
    rejected = client.post(
        "/sales", json={"product_id": 2, "quantity_sold": 1}, headers=api["staff"]
    )
    assert rejected.status_code == 400
    assert set(rejected.json) == {"error"}

    accepted = client.post(
        "/sales",
        json={"product_id": 2, "quantity_sold": 1, "buyer_id_proof": "ID-1234"},
        headers=api["staff"],
    )
    assert accepted.status_code == 201
    report = client.get("/reports/restricted-sales", headers=api["owner"])
    assert report.status_code == 200
    assert report.json["summary"]["with_buyer_id"] == 1
    assert report.json["records"][0]["has_buyer_id"] is True


def test_role_access_is_enforced(api):
    client = api["client"]
    body = {"name": "New Item"}

    assert client.post("/products", json=body, headers=api["staff"]).status_code == 403
    assert client.post("/products", json=body, headers=api["owner"]).status_code == 201


def test_insufficient_stock_does_not_change_inventory(api):
    response = api["client"].post(
        "/sales",
        json={"product_id": 1, "quantity_sold": 99},
        headers=api["staff"],
    )

    assert response.status_code == 400
    conn = sqlite3.connect(api["db_path"])
    remaining = conn.execute(
        "SELECT quantity_remaining FROM batches WHERE product_id = 1 ORDER BY expiry_date"
    ).fetchall()
    sale_count = conn.execute("SELECT COUNT(*) FROM sales").fetchone()[0]
    conn.close()
    assert remaining == [(3,), (4,)]
    assert sale_count == 0


def test_restock_suggestion_explains_season_and_current_stock(api):
    response = api["client"].get(
        "/suggestions/seasonal/1", headers=api["owner"]
    )

    assert response.status_code == 200
    assert response.json["current_stock"] == 7
    assert response.json["recommended_restock_quantity"] == 3
    assert "Basis:" in response.json["why"]
    assert "Season: Year-round" in response.json["why"]
    assert "Current usable stock is 7 units" in response.json["why"]


def test_export_routes_return_readable_xlsx_workbooks(api):
    client = api["client"]
    client.post(
        "/sales",
        json={"product_id": 2, "quantity_sold": 1, "buyer_id_proof": "ID-1234"},
        headers=api["staff"],
    )

    for path, expected_sheet in (
        ("/reports/inventory/export", "Inventory"),
        ("/reports/sales/export", "Sales"),
        ("/reports/compliance/export", "Compliance Log"),
    ):
        response = client.get(path, headers=api["owner"])
        assert response.status_code == 200
        assert response.mimetype == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        assert workbook.sheetnames == [expected_sheet]
        assert workbook.active.max_row >= 1
        workbook.close()
        assert client.get(path, headers=api["staff"]).status_code == 403


def test_all_documented_routes_respond(api):
    client = api["client"]
    headers = api["owner"]
    staff_headers = api["staff"]
    future = (date.today() + timedelta(days=365)).isoformat()

    checks = [
        client.get("/"),
        client.post(
            "/auth/login",
            json={"email": "owner@test.local", "password": "owner-password"},
        ),
        client.get("/auth/me", headers=headers),
        client.post(
            "/auth/register",
            json={
                "name": "Route Check",
                "email": "route-check@test.local",
                "password": "long-enough-password",
                "role": "staff",
            },
            headers=headers,
        ),
        client.get("/products", headers=staff_headers),
        client.post("/products", json={"name": "Route Check Product"}, headers=headers),
        client.get("/restricted-pesticides", headers=staff_headers),
        client.post(
            "/batches",
            json={
                "product_id": 1,
                "batch_no": "ROUTE-CHECK",
                "quantity_received": 2,
                "price_per_unit": 1.5,
                "purchase_date": date.today().isoformat(),
                "expiry_date": future,
            },
            headers=headers,
        ),
        client.get("/products/1/batches", headers=staff_headers),
        client.post(
            "/sales",
            json={"product_id": 1, "quantity_sold": 1},
            headers=staff_headers,
        ),
        client.get("/dashboard/summary", headers=headers),
        client.get("/dashboard/low-stock", headers=staff_headers),
        client.get("/dashboard/out-of-stock", headers=staff_headers),
        client.get("/dashboard/expiring?days=30", headers=staff_headers),
        client.get("/reports/restricted-sales", headers=headers),
        client.get("/forecast/history/1", headers=headers),
        client.get("/forecast/1", headers=headers),
        client.get("/forecast", headers=headers),
        client.get("/suggestions/seasonal/1", headers=headers),
        client.get("/suggestions/seasonal", headers=headers),
        client.get("/reports/inventory/export", headers=headers),
        client.get("/reports/sales/export", headers=headers),
        client.get("/reports/compliance/export", headers=headers),
    ]

    assert all(response.status_code in (200, 201) for response in checks)


def test_excel_export_escapes_formula_like_values():
    with backend.app.test_request_context("/"):
        response = backend.excel_response(
            "safe.xlsx", "Sheet", ["value"], [{"value": "=1+1"}]
        )
        response.direct_passthrough = False
        content = response.get_data()

    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.active["A2"].value == "'=1+1"
    workbook.close()


@pytest.mark.parametrize(
    "path,body",
    [
        ("/sales", {"product_id": 1, "quantity_sold": "many"}),
        ("/products", {"name": "Bad min", "min_stock_level": -1}),
        ("/batches", {"product_id": 1, "quantity_received": 2}),
    ],
)
def test_invalid_inputs_use_the_same_error_shape(api, path, body):
    response = api["client"].post(path, json=body, headers=api["owner"])

    assert response.status_code == 400
    assert set(response.json) == {"error"}


def test_invalid_query_parameters_use_the_same_error_shape(api):
    response = api["client"].get(
        "/dashboard/expiring?days=not-a-number", headers=api["staff"]
    )

    assert response.status_code == 400
    assert set(response.json) == {"error"}
