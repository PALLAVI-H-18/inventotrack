"""
exports.py - builds Excel (.xlsx) files in memory with openpyxl  (Oct 15)

Each function returns a BytesIO that Flask can send as a download.
Text cells are "neutralised" so a product name like  =HYPERLINK(...)  cannot run as a formula
when the owner opens the file (this is called formula injection).
"""
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import compliance
import stock

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
AMBER_FILL = PatternFill("solid", fgColor="FFE699")
RUPEE = '"₹"#,##0.00'


def safe(value):
    """Stop text that starts with = + - @ from being treated as an Excel formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def add_sheet(wb, title, headers, rows, money_columns=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([safe(v) for v in row])
    for col in money_columns:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = RUPEE
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, start=1):          # column width = longest value
        longest = max([len(str(header))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 45)
    return ws


def add_summary(wb, title, lines, first=False):
    ws = wb.active if first else wb.create_sheet(title=title, index=0)
    ws.title = title
    ws.append(["Item", "Value"])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for label, value in lines:
        ws.append([label, safe(value)])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28
    return ws


def to_buffer(wb):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generated_at():
    return datetime.now(stock.IST).strftime("%Y-%m-%d %H:%M IST")


def range_text(date_from, date_to):
    if not date_from and not date_to:
        return "All dates"
    return f"{date_from or 'start'} to {date_to or 'today'}"


# ---------------------------------------------------------------- inventory
def inventory_workbook(db):
    today = stock.today()
    today_s = today.strftime("%Y-%m-%d")
    soon = today + timedelta(days=30)

    products = stock.get_products_with_stock(db)
    batch_rows = db.execute("""
        SELECT p.name AS product_name, b.batch_no, s.name AS supplier_name, b.quantity_received,
               b.quantity_remaining, b.price_per_unit, b.purchase_date, b.expiry_date
        FROM batches b
        JOIN products p ON b.product_id = p.product_id
        LEFT JOIN suppliers s ON b.supplier_id = s.supplier_id
        ORDER BY p.name, b.expiry_date, b.batch_id
    """).fetchall()

    batches, value_ok, value_expired = [], 0.0, 0.0
    for b in batch_rows:
        days_left = (datetime.strptime(b["expiry_date"], "%Y-%m-%d").date() - today).days
        value = float(b["quantity_remaining"]) * float(b["price_per_unit"])
        if b["expiry_date"] < today_s:
            state = "Expired"
            value_expired += value
        else:
            state = "Expiring soon" if b["expiry_date"] <= soon.strftime("%Y-%m-%d") else "OK"
            value_ok += value
        batches.append([b["product_name"], b["batch_no"], b["supplier_name"] or "-", b["quantity_received"],
                        b["quantity_remaining"], float(b["price_per_unit"]), b["purchase_date"],
                        b["expiry_date"], days_left, state, round(value, 2)])

    wb = Workbook()
    add_summary(wb, "Summary", [
        ("Report", "Inventory"),
        ("Generated", generated_at()),
        ("Total products", len(products)),
        ("Out of stock", sum(1 for p in products if p["status"] == "out_of_stock")),
        ("Low stock", sum(1 for p in products if p["status"] == "low_stock")),
        ("Batches listed", len(batches)),
        ("Usable stock value (INR)", round(value_ok, 2)),
        ("Expired stock value (INR)", round(value_expired, 2)),
    ], first=True)

    add_sheet(wb, "Products",
              ["Product ID", "Name", "Brand", "Category", "Season", "Min stock", "Restricted", "Usable stock", "Status"],
              [[p["product_id"], p["name"], p["brand"] or "-", p["category"] or "-", p["season"] or "-",
                p["min_stock_level"], "Yes" if p["is_restricted"] else "No", p["total_quantity"],
                p["status"].replace("_", " ").title()] for p in products])

    ws = add_sheet(wb, "Batches",
                   ["Product", "Batch no", "Supplier", "Qty received", "Qty remaining", "Price per unit",
                    "Purchase date", "Expiry date", "Days left", "State", "Stock value"],
                   batches, money_columns=(6, 11))
    for row in ws.iter_rows(min_row=2):
        fill = RED_FILL if row[9].value == "Expired" else AMBER_FILL if row[9].value == "Expiring soon" else None
        if fill:
            for cell in row:
                cell.fill = fill
    return to_buffer(wb)


# ---------------------------------------------------------------- sales
def sales_workbook(db, date_from=None, date_to=None):
    query = """
        SELECT s.sale_id, s.sale_date, p.name AS product_name, b.batch_no, s.quantity_sold,
               b.price_per_unit, u.name AS sold_by_name, p.is_restricted, s.buyer_id_proof
        FROM sales s
        JOIN products p ON s.product_id = p.product_id
        JOIN batches b ON s.batch_id = b.batch_id
        LEFT JOIN users u ON s.sold_by = u.user_id
        WHERE 1 = 1
    """
    params = []
    if date_from:
        query += " AND s.sale_date >= ?"
        params.append(date_from)
    if date_to:
        query += " AND s.sale_date <= ?"
        params.append(date_to)
    query += " ORDER BY s.sale_date DESC, s.sale_id DESC"
    rows = db.execute(query, params).fetchall()

    data = [[r["sale_id"], r["sale_date"], r["product_name"], r["batch_no"], r["quantity_sold"],
             float(r["price_per_unit"]), r["sold_by_name"] or "-", "Yes" if r["is_restricted"] else "No",
             r["buyer_id_proof"] or "-"] for r in rows]

    wb = Workbook()
    add_summary(wb, "Summary", [
        ("Report", "Sales"),
        ("Generated", generated_at()),
        ("Date range", range_text(date_from, date_to)),
        ("Number of sale lines", len(data)),
        ("Units sold", sum(r["quantity_sold"] for r in rows)),
    ], first=True)
    add_sheet(wb, "Sales",
              ["Sale ID", "Date", "Product", "Batch no", "Qty sold", "Batch cost per unit", "Sold by",
               "Restricted", "Buyer ID"],
              data, money_columns=(6,))
    return to_buffer(wb)


# ---------------------------------------------------------------- compliance
def compliance_workbook(db, date_from=None, date_to=None, missing_only=False):
    summary, records = compliance.fetch_restricted_sales(db, date_from, date_to, missing_only)

    wb = Workbook()
    add_summary(wb, "Summary", [
        ("Report", "Restricted pesticide sales (compliance log)"),
        ("Generated", generated_at()),
        ("Date range", range_text(date_from, date_to)),
        ("Only sales missing a buyer ID", "Yes" if missing_only else "No"),
        ("Total records", summary["total_records"]),
        ("Total units", summary["total_units"]),
        ("With buyer ID", summary["with_buyer_id"]),
        ("Missing buyer ID", summary["missing_buyer_id"]),
    ], first=True)
    ws = add_sheet(wb, "Restricted sales",
                   ["Sale ID", "Date", "Product", "Batch no", "Qty sold", "Buyer ID", "Sold by", "Buyer ID status"],
                   [[r["sale_id"], r["sale_date"], r["product_name"], r["batch_no"], r["quantity_sold"],
                     (r["buyer_id_proof"] or "").strip() or "-", r["sold_by_name"] or "-",
                     "Recorded" if r["has_buyer_id"] else "MISSING"] for r in records])
    for row in ws.iter_rows(min_row=2):
        if row[7].value == "MISSING":
            for cell in row:
                cell.fill = RED_FILL
    return to_buffer(wb)
