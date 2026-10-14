"""Excel exports: the downloads must be real .xlsx files that open and contain the right data."""
from io import BytesIO

from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def open_xlsx(response):
    assert response.status_code == 200, response.get_data(as_text=True)[:200]
    assert response.mimetype == XLSX
    assert response.headers["Content-Disposition"].startswith("attachment; filename=")
    assert response.data[:2] == b"PK"                        # every .xlsx is a zip file
    return load_workbook(BytesIO(response.data))


def test_inventory_export(client, owner, world):
    wb = open_xlsx(client.get("/reports/inventory/export", headers=owner))
    assert wb.sheetnames == ["Summary", "Products", "Batches"]
    assert wb["Products"].max_row == 1 + 2                   # header + 2 products
    assert wb["Batches"].max_row == 1 + 4                    # header + 4 batches
    states = [row[9].value for row in wb["Batches"].iter_rows(min_row=2)]
    assert sorted(states) == ["Expired", "Expiring soon", "OK", "OK"]


def test_sales_export_and_date_filter(client, owner, staff, world):
    client.post("/sales", json={"product_id": world.neem, "quantity_sold": 15}, headers=staff)   # 2 sale lines
    wb = open_xlsx(client.get("/reports/sales/export", headers=owner))
    assert wb.sheetnames == ["Summary", "Sales"]
    assert wb["Sales"].max_row == 1 + 2
    empty = open_xlsx(client.get("/reports/sales/export?from=2999-01-01", headers=owner))
    assert empty["Sales"].max_row == 1                       # only the header, still a valid file


def test_compliance_export(client, owner, staff, world):
    client.post("/sales", json={"product_id": world.mono, "quantity_sold": 2, "buyer_id_proof": "AADHAAR-1234"}, headers=staff)
    wb = open_xlsx(client.get("/reports/compliance/export", headers=owner))
    assert wb.sheetnames == ["Summary", "Restricted sales"]
    row = [c.value for c in wb["Restricted sales"][2]]
    assert "Monocrotophos 36% SL" in row and "AADHAAR-1234" in row and "Recorded" in row


def test_formula_injection_is_neutralised(client, owner):
    client.post("/products", json={"name": '=HYPERLINK("http://evil.example","click")'}, headers=owner)
    wb = open_xlsx(client.get("/reports/inventory/export", headers=owner))
    names = [row[1].value for row in wb["Products"].iter_rows(min_row=2)]
    assert any(n.startswith("'=") for n in names)
    assert not any(n.startswith("=") for n in names)


def test_exports_are_owner_only(client, staff):
    for path in ["/reports/inventory/export", "/reports/sales/export", "/reports/compliance/export"]:
        assert client.get(path, headers=staff).status_code == 403


def test_export_with_bad_date_is_a_json_error(client, owner):
    response = client.get("/reports/sales/export?from=31-12-2026", headers=owner)
    assert response.status_code == 400 and "error" in response.get_json()
